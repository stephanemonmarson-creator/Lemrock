import openpyxl, csv, datetime, sys, collections
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)
def d(x):
    return x.strftime('%Y-%m-%d') if isinstance(x, datetime.datetime) else ''
def n(x):
    return round(x, 2) if isinstance(x, (int, float)) else ''
NET = {'Cdiscount': 'Cdiscount RM'}
TODAY = datetime.datetime(2026, 10, 14)

# --- campaigns
camps = []
for r in wb['Suivi Campagnes'].iter_rows(min_row=4, values_only=True):
    if not r[1] or not r[2] or str(r[1]).startswith(('TABLEAU', 'Mois', 'Campagnes', 'Budget', 'Le budget')):
        continue
    camps.append(r)
# --- clients
clients = {}
for r in wb['Suivi Client'].iter_rows(min_row=3, values_only=True):
    if not r[1] or r[1] in ('TOTAL',) or str(r[1]).startswith(('TABLEAU', 'Mois', 'Clients', 'Nouveaux', 'Taux', 'Croissance', 'Budget', 'Les 3', 'Total clients', 'Average')):
        continue
    if r[8] not in ('Live', 'Signed', 'Churn'):
        continue
    clients[r[1]] = r
csm = {}
for r in wb['CSM Client'].iter_rows(min_row=3, values_only=True):
    if r[1]:
        csm[r[1]] = r

# --- data from the S41 report (risk, tracker, next action, priority)
P1 = {'Boulanger','Monoprix','Bosch','Marshall','PV','Dorel','Cdiscount Mobile','Cdiscount Voyages','SNCF','Elhée','Qonto','DSR','Back Market'}
RISK = {**{k:'Red' for k in ['Boulanger','Monoprix','FidMe','Smartphone ID','Bosch','Marshall','PV','Archipel Kombucha','Basket4ballers','Allo Pneus','Mabex','Dorel']},
        'Cdiscount Mobile':'Orange',
        **{k:'Green' for k in ['Cdiscount Voyages','SNCF','Elhée','Qonto','Pour de Bon',"Levi's",'DSR']}}
TRACK = {'Elhée':'Installed','Qonto':'Installed','Archipel Kombucha':'Technical issue',
         **{k:'Installing' for k in ['Boulanger','Monoprix','Pour de Bon','Basket4ballers','Mabex']},
         **{k:'Proposed' for k in ['Cdiscount Voyages','Cdiscount Mobile','Allo Pneus',"Levi's"]},
         **{k:'Refused' for k in ['PV','DSR','SNCF','Bosch','Dorel','Marshall','FidMe','Smartphone ID']}}
REFUSAL = {'PV':'CJ refuses the integration','DSR':'Traffic-only objective','SNCF':'To re-propose end of year',
           'Bosch':'Orders sent manually by Cdiscount RM','Dorel':'Orders sent manually by Cdiscount RM','Marshall':'Orders sent manually by Cdiscount RM',
           'FidMe':'Consent issues (Adjust lead)','Smartphone ID':'Mobile app, wants a commercial incentive'}
NEXT = {'Boulanger':'Call André 14/10 11h30 – conversions file + tracker in prod',
        'Monoprix':'Clarify results per campaign, get renewal decision',
        'FidMe':'Fix WaterWipes offer, answer Sophie – stop 15/10 + review',
        'Smartphone ID':'Roxane answers Émile – explain clicks vs visits gap',
        'Bosch':'Summary per brand + hero products POC','Marshall':'Summary per brand + hero products POC',
        'PV':'Bookings via CJ, measurement solution, 1st perf review',
        'Archipel Kombucha':'Check consent and tracking with Clément',
        'Basket4ballers':'Remaining budget review, relaunch promo code',
        'Allo Pneus':'Review and end of campaign with Affilae','Mabex':'Call Julien – 15 min install with Clément',
        'Dorel':'Review and conditions for a restart','Cdiscount Mobile':'Call 15/10 15h30 – set profitability KPIs',
        'Cdiscount Voyages':'Highlight consolidated perf at 15/10 call','SNCF':'1st perf review, same date range',
        'Elhée':'Budget decision 19/10 16h','Qonto':'15-day review around 13/10','Pour de Bon':'Follow the new campaign',
        "Levi's":'Create Rakuten tracking links','DSR':'Review + chase Anthony for the Éducation nationale brief'}
# pending (signed, no campaign yet) – from S41 slides 7-8
PENDING = {
 'Total':('Lemrock','Contract to send; Total: offers, URLs, tag','Contract','To send'),
 'BUT':('Lemrock','Contract, Kbis/RIB, supplier form, URL cleaning','Contract','To send'),
 'DAZN':('Lemrock','Supplier form + quote to send, live 14–15/10','IO','To send'),
 'Shine':('Client','Signed, kick-off 15/10 16h; exit clause to settle','Contract','Signed'),
 'Peugeot':('Agency','Performics silent since 01/10, chased 05/10','',''),
 'Education Nationale':('Client','Brief end of October, launch early November','',''),
 'Puissante':('Client','Signed, kick-off 19/10 17h, live ~21–22/10','Contract','Signed'),
 'Jura':('Client','3 k€ or 5 k€/month, budget presented by CEO this week','',''),
 'Bonsoirs':('Client','Nothing new since mid-Sept, call 14/10 16h','',''),
 'Europ Assistance':('Client legal','Legal, compliance and procurement review from 12/10','Contract','Client legal review'),
 'Back Market':('Client legal','Contract at legal, chased 05/10 and 09/10','Contract','Client legal review'),
 'Balzac':('Client','Reduced to 1 month: product feed, objectives, call expected','',''),
 'TERACT':('Client','Files link expired, objectives per banner expected','',''),
 'VPG':('Client','Contract OK – stays feed, UTM, objectives expected','Contract','Signed'),
}
CHURN_MAP = {'Flux produits envoyé non tracké':'Untracked feed','Contact rompu':'Contact lost',
             'Essai non significatif':'Trial not significant'}

# --- networks per client (most frequent)
netc = collections.defaultdict(collections.Counter)
for r in camps:
    if r[7]: netc[r[1]][NET.get(r[7], r[7])] += 1

# --- companies
names = list(dict.fromkeys(list(clients) + [r[1] for r in camps]))
comp_rows = []
for c in names:
    k = csm.get(c)
    imp = k[4] if k else None
    comp_rows.append({
        'Company name': c,
        'Priority': 'P1' if (c in P1 or imp == 'VIP') else 'P2',
        'Network / Agency': netc[c].most_common(1)[0][0] if netc[c] else '',
        'Client type': '',
        'Flagship': '',
        'Risk level': RISK.get(c, ''),
        'Satisfaction score': '',
        'Conversion tracker': TRACK.get(c, ''),
        'Tracker refusal reason': REFUSAL.get(c, ''),
        'Next action': NEXT.get(c, ''),
    })

# --- deals
deal_rows = []
seen = collections.defaultdict(list)
camps.sort(key=lambda r: (r[1], r[16] or r[15] or datetime.datetime(2100,1,1)))
for r in camps:
    client, name, statut = r[1], r[2], r[3]
    start, end = r[16], r[17]
    if statut == 'Live': stage = 'Live'
    elif statut == 'Fin de campagne': stage = 'Ended'
    elif start and start <= TODAY: stage = 'Live'
    else: stage = 'Handover'
    prev = seen[client]
    amount = r[10] if isinstance(r[10], (int, float)) else 0
    if not prev: dtype = 'New business'
    elif amount > (prev[-1] or 0): dtype = 'Upsell'
    else: dtype = 'Renewal'
    seen[client].append(amount)
    nature = {'Performance':'Performance','Notoriété':'Awareness'}.get(r[4] or '', '')
    churn = r[5] or ''
    deal_rows.append({
        'Deal name': name, 'Company name': client, 'Pipeline': 'Campaigns', 'Deal stage': stage,
        'Deal type': dtype, 'Amount': n(r[10]), 'Total budget': n(r[12]),
        'Monthly budget Open Web': n(r[8]), 'Monthly budget ChatGPT': n(r[9]) or 0,
        'Duration (days)': n(r[11]), 'CPC': n(r[14]),
        'Signature date': d(r[15]), 'Close date': d(r[15]), 'Launch date': d(start), 'Campaign end date': d(end),
        'Network / Agency': NET.get(r[7], r[7] or ''), 'Inventory': {'ChatGPT':'ChatGPT','OpenWeb':'Open Web','Open Web':'Open Web','Tout Inventaire':'All inventory'}.get(r[22] or '', ''),
        'Objective type': nature, 'Objective source': r[23] or '', 'Objectives details': (r[24] or '').strip(),
        'Product feed': 'Catalog' if r[25] else 'HP' if r[26] else 'Category' if r[27] else 'Form' if r[28] else '',
        'Commitment document': '', 'Document status': '', 'Blocked by': '', 'Blocker comment': '',
        'Outcome': '', 'Churn reason': CHURN_MAP.get(churn, 'Other' if churn else ''), 'Churn details': ((churn + ' – ') if churn and churn not in CHURN_MAP else '') + (r[6] or ''),
    })
# outcome for ended deals: next deal of same client -> Renewed/Upsell, else Churn / Not renewed
by_client = collections.defaultdict(list)
for row in deal_rows: by_client[row['Company name']].append(row)
for c, rows in by_client.items():
    for i, row in enumerate(rows):
        if row['Deal stage'] != 'Ended': continue
        later = rows[i+1:]
        if later:
            row['Outcome'] = 'Upsell' if (later[0]['Amount'] or 0) > (row['Amount'] or 0) else 'Renewed'
        else:
            row['Outcome'] = 'Churn' if clients.get(c, [None]*9)[8] == 'Churn' else 'Not renewed'

# signed clients without any campaign
with_camp = {r[1] for r in camps}
for c, r in clients.items():
    if c in with_camp or r[8] != 'Signed': continue
    p = PENDING.get(c, ('', '', '', ''))
    lost = c == 'ManoMano'
    ow = r[4] if isinstance(r[4], (int, float)) else 0
    gpt = r[5] if isinstance(r[5], (int, float)) else 0
    dur = r[7] if isinstance(r[7], (int, float)) else ''
    deal_rows.append({
        'Deal name': f'{c} – Campaign 1', 'Company name': c, 'Pipeline': 'Campaigns',
        'Deal stage': 'Cancelled' if lost else 'Setup', 'Deal type': 'New business',
        'Amount': n(ow + gpt), 'Total budget': n((ow + gpt) * dur / 30) if dur else '',
        'Monthly budget Open Web': n(ow), 'Monthly budget ChatGPT': n(gpt), 'Duration (days)': n(dur), 'CPC': '',
        'Signature date': d(r[3]), 'Close date': d(r[3]), 'Launch date': '', 'Campaign end date': '',
        'Network / Agency': '', 'Inventory': '', 'Objective type': '', 'Objective source': '', 'Objectives details': r[22] or '',
        'Product feed': '', 'Commitment document': p[2], 'Document status': p[3], 'Blocked by': p[0], 'Blocker comment': p[1],
        'Outcome': '', 'Churn reason': 'Other' if lost else '', 'Churn details': 'Lost 06/10 (S41 report)' if lost else '',
    })

def write(path, rows):
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader(); w.writerows(rows)
write(f'{out}/companies.csv', comp_rows)
write(f'{out}/deals.csv', deal_rows)
T = {'Cdiscount Voyages', 'Boulanger', '1001 Pneus'}
write(f'{out}/TEST_companies.csv', [r for r in comp_rows if r['Company name'] in T])
write(f'{out}/TEST_deals.csv', [r for r in deal_rows if r['Company name'] in T])
print(len(comp_rows), 'companies', len(deal_rows), 'deals')
print(collections.Counter(r['Deal stage'] for r in deal_rows))
print('companies w/o deal:', [c for c in names if c not in by_client and c not in clients])
