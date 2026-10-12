from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
O=RGBColor(0xEE,0x6A,0x2C);K=RGBColor(0x11,0x11,0x11);G=RGBColor(0x55,0x5B,0x66);LG=RGBColor(0x8A,0x93,0xA6)
GR=RGBColor(0x1E,0x7B,0x34);RD=RGBColor(0xB0,0x48,0x37);BG=RGBColor(0xF8,0xF6,0xF2);W=RGBColor(255,255,255);BOX=RGBColor(0xF1,0xEF,0xEA)
import json,sys
D=json.load(open(sys.argv[1]))
p=Presentation();p.slide_width=Emu(12192000);p.slide_height=Emu(6858000)
s=p.slides.add_slide(p.slide_layouts[6])
X=lambda v:Emu(int(v*12192000/2000))  # coords in 2000-wide px space
def rect(x,y,w,h,fill,line=None):
    r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,X(x),X(y),X(w),X(h));r.fill.solid();r.fill.fore_color.rgb=fill
    if line: r.line.color.rgb=line;r.line.width=Pt(.75)
    else: r.line.fill.background()
    r.shadow.inherit=False
    from pptx.oxml.ns import qn
    sp=r._element.spPr
    for e in sp.findall(qn('a:effectLst')): sp.remove(e)
    from lxml import etree; sp.append(etree.SubElement(sp,qn('a:effectLst')))
    st=r._element.find(qn('p:style'))
    if st is not None: r._element.remove(st)
    return r
def txt(x,y,w,h,runs,size=12,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    tb=s.shapes.add_textbox(X(x),X(y),X(w),X(h));tf=tb.text_frame;tf.word_wrap=True;tf.vertical_anchor=anchor
    for m in ('margin_left','margin_right','margin_top','margin_bottom'):setattr(tf,m,0)
    lines=runs if isinstance(runs[0],list) else [runs]
    for i,line in enumerate(lines):
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.alignment=align
        for t,c,sz,b,it in line:
            r=para.add_run();r.text=t;f=r.font;f.size=Pt(sz);f.bold=b;f.italic=it;f.color.rgb=c;f.name='Arial'
    return tb
TR={'RAPPORT DE PERFORMANCE · 01/09/2026 → 30/09/2026':'PERFORMANCE REPORT · 01/09/2026 → 30/09/2026','COMPARAISON':'COMPARISON','DIFFUSION':'REACH','vs période préc.':'vs prev. period',
'INSIGHTS QUALITATIFS · ANALYSE DES RECHERCHES':'QUALITATIVE INSIGHTS · SEARCH ANALYSIS','CHAMP LEXICAL':'LEXICAL FIELD','Termes les plus fréquents':'Most frequent terms','EXEMPLES DE REQUÊTES':'SAMPLE QUERIES',
'Questions posées aux IA':'Questions asked to AI','CATÉGORIES DE PRODUITS':'PRODUCT CATEGORIES','Répartition des requêtes':'Campaign product mix','INSIGHTS QUALITATIFS · SEPTEMBRE 2026':'QUALITATIVE INSIGHTS · SEPTEMBER 2026',
'OBJECTIF   ':'OBJECTIVE   ','INSIGHTS QUALITATIFS':'QUALITATIVE INSIGHTS','À définir':'To be defined','CLICS':'CLICKS','N/C':'N/A'}
R=lambda t,c=K,sz=12,b=False,it=False:((TR.get(t,t) if D.get('lang')=='en' else t),c,sz,b,it)
rect(0,0,2000,1125,BG)
txt(87,80,1200,30,[R('RAPPORT DE PERFORMANCE · 01/09/2026 → 30/09/2026',O,13,True)])
txt(87,105,1200,80,[R(D['client'],K,36,True)])
txt(1300,97,542,25,[R('COMPARAISON',LG,12,True)],align=PP_ALIGN.RIGHT)
txt(1300,128,542,30,[R('vs 01/08/2026 → 31/08/2026',K,15,True)],align=PP_ALIGN.RIGHT)
for x,t in((87,'DIFFUSION'),(392,'ENGAGEMENT'),(1002,'CONVERSION')):txt(x,210,400,30,[R(t,O,14,True)])
kp=[tuple(k[:3])+(GR if k[3]=='up' else RD if k[3]=='down' else G,k[4]) for k in D['kpis']]
for i,(l,v,d,dc,y) in enumerate(kp):
    x=87+i*305;gr=i>=3
    if gr:
        rect(x,252,228,185,RGBColor(0xF1,0xEF,0xEB))
        pl=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,X(x+118),X(238),X(100),X(28));pl.adjustments[0]=0.5;pl.fill.solid();pl.fill.fore_color.rgb=RGBColor(0xD9,0xD6,0xD1);pl.line.fill.background();pl.shadow.inherit=False
        txt(x+118,238,100,28,[R('INDICATIF',G,8,True)],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        dc=LG;vc=LG
    else:
        rect(x,252,228,185,W);rect(x,252,4,185,O);vc=K
    txt(x+22,280 if gr else 265,200,20,[R(l,LG if gr else G,10,True)]);txt(x+22,310,200,45,[R(v,vc,22 if len(v)<9 else 18,True)])
    txt(x+22,368,200,18,[R(d,dc,10,True)]);txt(x+22,410,200,15,[R(y,LG if gr else G,8,False,True)])
if D.get('conv_msg'):
    tri=s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,X(1405),X(441),X(30),X(14));tri.fill.solid();tri.fill.fore_color.rgb=RGBColor(0xEB,0xE9,0xE5);tri.line.fill.background();tri.shadow.inherit=False
    b=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,X(1002),X(455),X(838),X(72));b.adjustments[0]=0.15;b.fill.solid();b.fill.fore_color.rgb=RGBColor(0xEB,0xE9,0xE5);b.line.fill.background();b.shadow.inherit=False
    m=D['conv_msg'];h,_,t=m.partition(' : ') if ' : ' in m else ('',None,m)
    txt(1022,462,798,60,[R(h+' : ',G,9,True),R(t,G,9)] if h else [R(t,G,9)],anchor=MSO_ANCHOR.MIDDLE)
txt(87,545,1200,30,[R('INSIGHTS QUALITATIFS · ANALYSE DES RECHERCHES',O,14,True)])
for x in(87,688,1289):rect(x,588,558,420,W)
# champ lexical
txt(127,618,480,25,[R('CHAMP LEXICAL',G,13,True)]);txt(127,655,480,20,[R('Termes les plus fréquents',LG,10,False,True)])
words=[[R(t,{'O':O,'K':K,'G':K}[c],sz,bold) for t,c,sz,bold in line] for line in D['words']]
tb=txt(127,695,490,300,words)
for para in tb.text_frame.paragraphs: para.space_after=Pt(5)
for i in range(1,len(words)):pass
# requetes
txt(728,618,480,25,[R('EXEMPLES DE REQUÊTES',G,13,True)]);txt(728,655,480,20,[R('Questions posées aux IA',LG,10,False,True)])
qs=D['queries']
for i,q in enumerate(qs):
    y=700+i*60;txt(728,y,478,50,[R(q,K,10)],anchor=MSO_ANCHOR.MIDDLE)
    if i<len(qs)-1: rect(728,y+55,478,1,RGBColor(0xE6,0xE3,0xDE))
# categories
txt(1329,618,480,25,[R(D.get('cats_title','CATÉGORIES DE PRODUITS'),G,13,True)]);txt(1329,655,480,20,[R(D.get('cats_sub','Répartition des requêtes'),LG,10,False,True)])
PAL=[RGBColor(0x1C,0x7C,0x8C),RGBColor(0x7A,0x4E,0x96),RGBColor(0xEB,0x9A,0x5E),O,RGBColor(0x3E,0x5A,0x78),RGBColor(0xA0,0x92,0x78),RGBColor(0x8C,0xB3,0x6B),RGBColor(0xC9,0x5B,0x7A)]
cats=[(c[0],c[1],c[2] if len(c)>2 else f'{c[1]} %') for c in D['cats']]
n=len(cats);rh=min(52,300/n);TRK=RGBColor(0xEC,0xEA,0xE5);BW=330;mx=max(c[1] for c in cats)
for i,(nm,v,lab) in enumerate(cats):
    yy=700+i*rh
    txt(1329,yy,360,22,[R(nm,K,10 if len(nm)<38 else 9)])
    rect(1329,yy+25,BW,9,TRK)
    if v>0: rect(1329,yy+25,max(4,BW*v/100),9,O if v==mx else RGBColor(0x9A,0xA3,0xB2))
    txt(1680,yy+12,130,30,[R(lab,K if v>=1 else LG,13 if v>=1 else 10,True)],align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)
s=p.slides.add_slide(p.slide_layouts[6])
rect(0,0,2000,1125,BG)
txt(95,70,1200,30,[R('INSIGHTS QUALITATIFS · SEPTEMBRE 2026',O,13,True)])
txt(95,100,1200,80,[R(D['client'],K,36,True)])
txt(1500,82,310,30,[R('Lemrock',K,15,True)],align=PP_ALIGN.RIGHT)
txt(95,200,1500,30,[R('OBJECTIF   ',O,13,True),R(D.get('objectif','À définir'),K,13)])
txt(95,250,800,30,[R('INSIGHTS QUALITATIFS',G,14,True)])
ins=D['insights']
for i,(t,b,r) in enumerate(ins):
    x=95+(i%2)*888;y=298+(i//2)*380
    rect(x,y,832,348,W)
    txt(x+42,y+35,80,40,[R(f'0{i+1}',O,20,True)])
    txt(x+132,y+38,660,60,[R(t,K,14,True)])
    txt(x+42,y+130,750,120,[R(b,G,11)])
    txt(x+42,y+245,750,80,[R('→  '+r,G,9,True)])

from pptx.oxml.ns import qn as _qn
for _sl in p.slides:
    for _sh in _sl.shapes:
        _st=_sh._element.find(_qn('p:style'))
        if _st is not None: _sh._element.remove(_st)
p.save(D['out'])
